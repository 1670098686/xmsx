# -*- coding: utf-8 -*-
"""表格结构化能力（extract_tables）的单元测试（无网络、无数据库）。

覆盖中间件对 PDF 线框表 / DOCX 内嵌表 / 多表同页的解析完整性：
- tool_layer.file_parser：PyMuPDF 绘制的 ruled 中文表经 pdfplumber 识别为
  GitHub 风格 Markdown；python-docx 表格直抽；TXT 与旧版 DOC 返回空；
- tool_layer.document_vision：体检抽样页 table_pages 线框探测；
- agent.tools：extract_tables 已注册进工具表并可直接调用；
- agent.nodes.observe_node：表格以「【表格数据·第N页】」/「【Word表N】」
  块并入全文末尾，同页多表合并为单块，重复提取不产生重复块、按键去重。

造中文 PDF 必须 fontname="china-s"；import pymupdf 前先导入
tool_layer.document_vision，让内置 vendor 版优先于本机全局安装。
"""
import pytest

# 先导入 document_vision 再导入 pymupdf（vendor 内置版优先，测试环境约定）
from tool_layer import document_vision, file_parser  # noqa: E402
from agent.nodes.observe_node import observe_node  # noqa: E402
from agent.state import TOOL_EXTRACT_TABLES  # noqa: E402
from agent.tools import build_tool_registry  # noqa: E402
from config import constants as C  # noqa: E402

import pymupdf  # noqa: E402

# 单表：3 行 2 列（表头 + 2 行数据）
_TABLE_ROWS = [
    ("指标", "数值"),
    ("准确率", "0.95"),
    ("召回率", "0.88"),
]
# 同页第二张表：2 行 2 列
_TABLE_ROWS_2 = [
    ("模块", "得分"),
    ("解析模块", "92"),
]


def _draw_ruled_table(page, top, rows, x0=60, x_mid=220, x2=380, row_h=30):
    """在 PDF 页指定纵向位置绘制一张带完整线框的中文表。

    Args:
        page: PyMuPDF 页对象。
        top: 表格上边界 y 坐标。
        rows: 行数据（首行为表头）。
        x0/x_mid/x2: 两条列分隔的三条竖线 x 坐标。
        row_h: 行高（pt）。
    """
    ys = [top + index * row_h for index in range(len(rows) + 1)]
    for x in (x0, x_mid, x2):
        page.draw_line((x, ys[0]), (x, ys[-1]),
                       color=(0, 0, 0), width=0.8)
    for y in ys:
        page.draw_line((x0, y), (x2, y), color=(0, 0, 0), width=0.8)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            page.insert_text(((x0, x_mid)[c] + 8, ys[r] + 20),
                             text, fontname="china-s", fontsize=10)


def _build_pdf_with_tables(path, table_groups):
    """构造含一张或多张线框表的中文 PDF。

    Args:
        path: 输出 PDF 路径。
        table_groups: 每张表的行数据列表（纵向排列在同一页）。
    """
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 80), "表1 模型评测指标结果对比表（含表格数据）",
                     fontname="china-s", fontsize=12)
    top = 110
    for rows in table_groups:
        _draw_ruled_table(page, top, rows)
        top += 30 * (len(rows) + 1) + 40
    doc.save(path)
    doc.close()


def _build_plain_text_pdf(path):
    """无表格线框的普通中文文本 PDF（用于表格探测阴性对照）。"""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(
        pymupdf.Rect(50, 50, 545, 800),
        "本文研究模块化设计方法，通过多案例研究对比多个系统的迭代效率，"
        "结论表明模块化能够显著降低维护成本并缩短交付周期。",
        fontname="china-s", fontsize=11)
    doc.save(path)
    doc.close()


@pytest.fixture
def ruled_table_pdf(tmp_path):
    path = tmp_path / "ruled_table.pdf"
    _build_pdf_with_tables(str(path), [_TABLE_ROWS])
    return str(path)


@pytest.fixture
def two_tables_pdf(tmp_path):
    path = tmp_path / "two_tables.pdf"
    _build_pdf_with_tables(str(path), [_TABLE_ROWS, _TABLE_ROWS_2])
    return str(path)


@pytest.fixture
def plain_text_pdf(tmp_path):
    path = tmp_path / "plain.pdf"
    _build_plain_text_pdf(str(path))
    return str(path)


@pytest.fixture
def table_docx(tmp_path):
    import docx

    path = tmp_path / "table.docx"
    document = docx.Document()
    document.add_paragraph("实验数据如下表所示。")
    table = document.add_table(rows=3, cols=2)
    cells = [["指标", "数值"], ["准确率", "0.95"], ["召回率", "0.88"]]
    for r in range(3):
        for c in range(2):
            table.rows[r].cells[c].text = cells[r][c]
    document.save(str(path))
    return str(path)


@pytest.fixture
def plain_txt(tmp_path):
    path = tmp_path / "plain.txt"
    path.write_text("纯文本文献，没有任何表格结构。" * 3, encoding="utf-8")
    return str(path)


# ================= file_parser：表格 → Markdown =================

def test_pdf_ruled_table_extracted_as_markdown(ruled_table_pdf):
    result = file_parser.extract_tables_markdown(ruled_table_pdf, 0, 20)
    assert result["table_pages"] == [0]
    assert len(result["tables"]) == 1
    table = result["tables"][0]
    assert table["page"] == 0 and table["index"] == 0
    lines = table["markdown"].splitlines()
    # GitHub 风格：首行表头、第二行分隔、其后数据行
    assert lines[0] == "| 指标 | 数值 |"
    assert lines[1] == "| --- | --- |"
    assert lines[2:] == ["| 准确率 | 0.95 |", "| 召回率 | 0.88 |"]


def test_pdf_extract_all_pages_respects_limit(two_tables_pdf):
    all_tables = file_parser.extract_tables_markdown(two_tables_pdf, -1, 20)
    assert [t["index"] for t in all_tables["tables"]] == [0, 1]
    assert all_tables["table_pages"] == [0]  # 两表同页，页码只记一次

    limited = file_parser.extract_tables_markdown(two_tables_pdf, -1, 1)
    assert len(limited["tables"]) == 1


def test_docx_table_extracted_as_markdown(table_docx):
    result = file_parser.extract_tables_markdown(table_docx)
    assert result["table_pages"] == []
    assert len(result["tables"]) == 1
    table = result["tables"][0]
    assert table["page"] == "docx" and table["index"] == 0
    markdown = table["markdown"]
    assert "| 指标 | 数值 |" in markdown
    assert "| --- | --- |" in markdown
    assert "| 准确率 | 0.95 |" in markdown


def test_txt_has_no_tables(plain_txt):
    result = file_parser.extract_tables_markdown(plain_txt)
    assert result == {"tables": [], "table_pages": []}


# ================= document_vision：体检表格页探测 =================

def test_inspect_layout_detects_table_pages(ruled_table_pdf):
    layout = document_vision.inspect_document_layout(ruled_table_pdf)
    assert layout["table_pages"] == [0]


def test_inspect_layout_plain_pdf_has_no_table_pages(plain_text_pdf):
    layout = document_vision.inspect_document_layout(plain_text_pdf)
    assert layout["table_pages"] == []


# ================= Agent 工具注册与 Observe 合并 =================

def test_extract_tables_tool_registered_and_invokable(ruled_table_pdf):
    registry = build_tool_registry(None)
    assert TOOL_EXTRACT_TABLES in registry
    out = registry[TOOL_EXTRACT_TABLES].invoke(
        {"file_path": ruled_table_pdf, "page_index": 0, "max_tables": 20})
    assert out["ok"] is True
    assert out["table_pages"] == [0]
    assert "准确率" in out["tables"][0]["markdown"]


def _observe_table_output(state, output):
    """把一次 extract_tables 工具输出喂给 Observe 节点并回写增量。

    Args:
        state: Agent 共享状态（会被原地更新）。
        output: extract_tables 工具的返回字典。
    Returns:
        Observe 节点的状态增量。
    """
    state["tool_history"] = [{
        "step": int(state.get("current_step", 0)) + 1,
        "tool": TOOL_EXTRACT_TABLES,
        "args": {"page_index": -1, "max_tables": 20},
        "output": output,
    }]
    state["current_step"] = state["tool_history"][-1]["step"]
    update = observe_node(state)
    state.update(update)
    return update


def test_observe_appends_pdf_table_block_to_abs_text(ruled_table_pdf):
    out = build_tool_registry(None)[TOOL_EXTRACT_TABLES].invoke(
        {"file_path": ruled_table_pdf, "page_index": -1, "max_tables": 20})
    state = {"abs_text": "文献正文内容。", "table_texts": {},
             "tool_history": []}
    update = _observe_table_output(state, out)
    assert "table_texts" in update and "abs_text" in update
    assert state["abs_text"].startswith("文献正文内容。")
    assert state["abs_text"].rstrip().endswith("| 召回率 | 0.88 |")
    assert "【表格数据·第1页】" in state["abs_text"]


def test_observe_merges_same_page_tables_into_single_block(two_tables_pdf):
    out = build_tool_registry(None)[TOOL_EXTRACT_TABLES].invoke(
        {"file_path": two_tables_pdf, "page_index": -1, "max_tables": 20})
    state = {"abs_text": "正文。", "table_texts": {}, "tool_history": []}
    _observe_table_output(state, out)
    # 同页两表合并进一个页块，但两张表的单元格内容都在
    assert state["abs_text"].count("【表格数据·第1页】") == 1
    assert "准确率" in state["abs_text"]
    assert "解析模块" in state["abs_text"]
    assert len(state["table_texts"]) == 2  # 仍按 (页, 表序) 保留两个去重键


def test_observe_table_merge_is_idempotent(ruled_table_pdf):
    """重复提取同一表格：全文块不重复追加，table_texts 按键去重。"""
    registry = build_tool_registry(None)
    state = {"abs_text": "正文。", "table_texts": {}, "tool_history": []}
    for _ in range(2):
        out = registry[TOOL_EXTRACT_TABLES].invoke(
            {"file_path": ruled_table_pdf, "page_index": -1, "max_tables": 20})
        _observe_table_output(state, out)
    assert state["abs_text"].count("【表格数据·第1页】") == 1
    assert state["abs_text"].count("| 召回率 | 0.88 |") == 1
    assert len(state["table_texts"]) == 1


def test_observe_docx_table_uses_word_block_title(table_docx):
    out = build_tool_registry(None)[TOOL_EXTRACT_TABLES].invoke(
        {"file_path": table_docx, "page_index": -1, "max_tables": 20})
    state = {"abs_text": "正文。", "table_texts": {}, "tool_history": []}
    _observe_table_output(state, out)
    assert "【Word表1】" in state["abs_text"]
    assert "准确率" in state["abs_text"]
    assert C.AGENT_TABLE_BLOCK_TITLE not in state["abs_text"].split("【Word表1】")[0]
