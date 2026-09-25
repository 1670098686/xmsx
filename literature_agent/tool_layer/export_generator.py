"""导出文件生成器：解析报告导出为 Word / TXT / PDF。

纯文件生成工具，不访问数据库；报告数据由 business 层组装传入。
"""
import os

from utils.exceptions import FileInvalidError, PermissionDeniedError
from utils.logger import get_logger

logger = get_logger()

# 报告字段 → 章节标题（与 ParsePage 7 个折叠块一致）
SECTION_TITLES = [
    ("research_background", "一、研究背景"),
    ("core_view", "二、核心观点"),
    ("research_method", "三、研究方法"),
    ("innovation_point", "四、创新点"),
    ("research_conclusion", "五、研究结论"),
    ("reference_list", "六、参考文献"),
]

EXPORT_FMT_WORD = "word"
EXPORT_FMT_TXT = "txt"
EXPORT_FMT_PDF = "pdf"
SUPPORTED_EXPORT_FMT = (EXPORT_FMT_WORD, EXPORT_FMT_TXT, EXPORT_FMT_PDF)


def build_section_lines(report_data: dict) -> list:
    """把报告字典整理为 [(标题, 正文), ...] 有序序列。

    Args:
        report_data: 含文献元信息与 6 个结构字段的报告字典。
    Returns:
        [(section_title, content), ...]，自动跳过空章节。
    """
    sections = []
    title = report_data.get("literature_title") or report_data.get("title") or "未命名文献"
    author = report_data.get("literature_author", "")
    publish_time = report_data.get("publish_time", "")
    journal_source = report_data.get("journal_source", "")
    keywords = report_data.get("keywords", "")

    basic_lines = [
        f"文献标题：{title}",
        f"作者：{author or '未知'}",
        f"发表时间：{publish_time or '未知'}",
        f"期刊/来源：{journal_source or '未知'}",
    ]
    if keywords:
        if isinstance(keywords, (list, tuple)):
            keywords = "、".join(str(k) for k in keywords)
        basic_lines.append(f"关键词：{keywords}")
    sections.append(("文献基础信息", "\n".join(basic_lines)))

    for field, section_title in SECTION_TITLES:
        content = (report_data.get(field) or "").strip()
        if content:
            sections.append((section_title, content))
    return sections


def export_to_txt(report_data: dict, output_path: str) -> None:
    """生成 TXT 报告。

    Args:
        report_data: 报告数据。
        output_path: 目标文件路径（.txt）。
    Raises:
        PermissionDeniedError: 目录不可写。
    """
    _ensure_writable(output_path, ".txt")
    lines = []
    for title, content in build_section_lines(report_data):
        lines.append(f"{title}\n{'=' * 40}\n{content}\n")
    try:
        with open(output_path, "w", encoding="utf-8") as fp:
            fp.write("\n".join(lines))
    except OSError as exc:
        raise PermissionDeniedError(f"TXT 导出失败：{exc}") from exc
    logger.info("TXT 报告已导出：%s", output_path)


def export_to_word(report_data: dict, output_path: str) -> None:
    """生成 Word(.docx) 报告。

    Args:
        report_data: 报告数据。
        output_path: 目标文件路径（.docx）。
    Raises:
        FileInvalidError: python-docx 不可用。
    """
    _ensure_writable(output_path, ".docx")
    try:
        import docx
        from docx.shared import Pt
    except ImportError as exc:
        raise FileInvalidError("缺少 python-docx 依赖，无法导出 Word") from exc

    document = docx.Document()
    document.add_heading(
        report_data.get("literature_title") or "文献解析报告", level=1
    )
    for title, content in build_section_lines(report_data):
        document.add_heading(title, level=2)
        for paragraph_text in content.split("\n"):
            paragraph = document.add_paragraph(paragraph_text)
            for run in paragraph.runs:
                run.font.size = Pt(11)
    try:
        document.save(output_path)
    except OSError as exc:
        raise PermissionDeniedError(f"Word 导出失败：{exc}") from exc
    logger.info("Word 报告已导出：%s", output_path)


def export_to_pdf(report_data: dict, output_path: str) -> None:
    """生成 PDF 报告（reportlab + 内置 STSong-Light 中文字体）。

    Args:
        report_data: 报告数据。
        output_path: 目标文件路径（.pdf）。
    Raises:
        FileInvalidError: reportlab 不可用。
    """
    _ensure_writable(output_path, ".pdf")
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise FileInvalidError("缺少 reportlab 依赖，无法导出 PDF") from exc

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except Exception as exc:  # 字体重复注册等情况忽略
        logger.warning("PDF 中文字体注册异常（可忽略）：%s", exc)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CnTitle", parent=styles["Title"], fontName="STSong-Light", fontSize=18
    )
    heading_style = ParagraphStyle(
        "CnHeading", parent=styles["Heading2"], fontName="STSong-Light", fontSize=14
    )
    body_style = ParagraphStyle(
        "CnBody", parent=styles["BodyText"], fontName="STSong-Light",
        fontSize=11, leading=18,
    )

    story = [Paragraph(report_data.get("literature_title") or "文献解析报告", title_style)]
    for title, content in build_section_lines(report_data):
        story.append(Spacer(1, 8))
        story.append(Paragraph(title, heading_style))
        for paragraph_text in content.split("\n"):
            if paragraph_text.strip():
                safe_text = paragraph_text.replace("&", "&amp;").replace("<", "&lt;")
                story.append(Paragraph(safe_text, body_style))
    try:
        SimpleDocTemplate(output_path, pagesize=A4).build(story)
    except OSError as exc:
        raise PermissionDeniedError(f"PDF 导出失败：{exc}") from exc
    logger.info("PDF 报告已导出：%s", output_path)


def _ensure_writable(output_path: str, expected_suffix: str) -> None:
    """导出路径公共校验。

    Args:
        output_path: 目标路径。
        expected_suffix: 期望后缀（.txt/.docx/.pdf）。
    Raises:
        FileInvalidError: 路径为空或后缀不符。
        PermissionDeniedError: 目录不可写。
    """
    if not output_path:
        raise FileInvalidError("导出路径为空")
    if not output_path.lower().endswith(expected_suffix):
        raise FileInvalidError(f"导出文件后缀必须为 {expected_suffix}")
    output_dir = os.path.dirname(os.path.abspath(output_path))
    try:
        os.makedirs(output_dir, exist_ok=True)
    except OSError as exc:
        raise PermissionDeniedError(f"导出目录不可用：{exc}") from exc


def export_report(report_data: dict, output_path: str, fmt: str) -> str:
    """按格式分发导出。

    Args:
        report_data: 报告数据。
        output_path: 目标文件路径。
        fmt: word/txt/pdf。
    Returns:
        实际写入的文件路径。
    Raises:
        FileInvalidError: 不支持的导出格式。
    """
    fmt = (fmt or "").lower()
    if fmt == EXPORT_FMT_WORD:
        export_to_word(report_data, output_path)
    elif fmt == EXPORT_FMT_TXT:
        export_to_txt(report_data, output_path)
    elif fmt == EXPORT_FMT_PDF:
        export_to_pdf(report_data, output_path)
    else:
        raise FileInvalidError(f"不支持的导出格式：{fmt}（仅支持 word/txt/pdf）")
    return output_path
